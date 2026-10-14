#!/usr/bin/env python3
"""
Llena el libro "Control_Variables_Produccion.xlsx" a partir de un repositorio LOCAL.
Puede recorrer TODAS las aplicaciones de un monorepo (una por subcarpeta).

Qué hace (todo en tu computadora, sin conexiones a internet):
  1. Lee los .properties y los .env de cada aplicación.
  2. Guarda SOLO nombres de variables, archivo, ambiente, si es sensible y el commit de Git.
  3. Genera registro_variables.csv y boveda_secretos.csv.
  4. Si tienes openpyxl y le pasas el libro, llena las pestañas y guarda una copia nueva.

Lo que NUNCA hace:
  - Escribir valores de variables en ningún archivo ni en pantalla.
  - Modificar el repositorio.
  - Usar la red.

Uso (Windows, PowerShell):
  Todas las aplicaciones de un monorepo (una subcarpeta por aplicación):
    py llenar_control_variables.py --repo "C:\\ruta\\APC_Liberaciones_2026" --todas --xlsx "Control_Variables_Produccion.xlsx"

  Una sola aplicación:
    py llenar_control_variables.py --repo "C:\\ruta\\APC_Liberaciones_2026\\asp-pago-onb" --xlsx "Control_Variables_Produccion.xlsx"

  Solo CSV (sin openpyxl):
    py llenar_control_variables.py --repo "C:\\ruta\\APC_Liberaciones_2026" --todas --solo-csv

Opcional:
  --responsable "Nombre"       responsable de cada fila (por defecto [pendiente])
  --salida "salida_variables"  carpeta donde van los CSV y el libro llenado
"""
import argparse
import csv
import re
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

AMBIENTES = {
    "dev": "Desarrollo",
    "test": "QA",
    "qa": "QA",
    "pre": "Preproducción",
    "prod": "Producción",
    "local": "Local",
}
# Nombres de llave que sugieren un secreto. Revisar a mano: marca de más y de menos.
SENSIBLE_LLAVE = re.compile(
    r"(pass|pwd|secret|token|api[_-]?key|clave|credential|private|keystore|auth)", re.IGNORECASE
)
# Valor con usuario y contraseña embebidos en una URL
URL_CON_CREDENCIALES = re.compile(r"://[^/\s:@]+:[^/\s@]+@")
EXCLUIR_CARPETAS = {".git", "target", "node_modules", ".idea", ".vscode"}
FILA_INICIAL = 6  # primera fila de captura en las pestañas del libro
ULTIMA_FILA_FORMULAS = 50000


def clave_de_linea(linea):
    s = linea.lstrip()
    if not s or s[0] in "#!":
        return None, None
    m = re.match(r"([^\s=:]+)\s*[=:]?\s*(.*)$", s)
    if not m:
        return None, None
    return m.group(1), m.group(2)


def leer_llaves(ruta):
    """(llave, valor). El valor solo vive en memoria para detectar credenciales; nunca se escribe."""
    pares = []
    continuacion = False
    try:
        texto = ruta.read_text(encoding="utf-8", errors="replace")
    except OSError as e:
        print(f"  [aviso] no se pudo leer {ruta.name}: {e}", file=sys.stderr)
        return pares
    for linea in texto.splitlines():
        if continuacion:
            continuacion = linea.rstrip().endswith("\\")
            continue
        clave, valor = clave_de_linea(linea)
        if clave:
            pares.append((clave, valor))
            continuacion = valor.rstrip().endswith("\\")
    return pares


def es_sensible(clave, valor):
    return bool(SENSIBLE_LLAVE.search(clave)) or bool(URL_CON_CREDENCIALES.search(valor))


def ambiente_de(ruta):
    tokens = re.split(r"[-.]", ruta.name.lower())
    for t in reversed(tokens):
        if t in AMBIENTES:
            return AMBIENTES[t], t
    return "Todos (base)", "base"


def base_de(ruta):
    partes = [p for p in ruta.stem.split("-") if p.lower() not in AMBIENTES]
    return "-".join(partes) or ruta.stem


def sha_git(carpeta):
    try:
        r = subprocess.run(
            ["git", "-C", str(carpeta), "rev-parse", "--short", "HEAD"],
            capture_output=True, text=True, timeout=15,
        )
        if r.returncode == 0 and r.stdout.strip():
            return r.stdout.strip()
    except (OSError, subprocess.SubprocessError):
        pass
    return "[pendiente]"


def buscar_archivos(app_dir):
    archivos = []
    for ruta in app_dir.rglob("*"):
        if any(p in EXCLUIR_CARPETAS for p in ruta.relative_to(app_dir).parts):
            continue
        if not ruta.is_file():
            continue
        nombre = ruta.name.lower()
        if nombre.endswith(".properties") or (ruta.parent == app_dir and nombre.startswith(".env")):
            archivos.append(ruta)
    return sorted(archivos)


def descubrir_apps(raiz):
    """Subcarpetas del monorepo que contienen .properties o .env (una por aplicación)."""
    apps = []
    for carpeta in sorted(raiz.iterdir()):
        if not carpeta.is_dir() or carpeta.name.startswith(".") or carpeta.name in EXCLUIR_CARPETAS:
            continue
        if buscar_archivos(carpeta):
            apps.append(carpeta)
    return apps


def extraer_app(raiz, app_dir, app, responsable, id_inicio):
    """Devuelve (registro, boveda, diferencias) de una aplicación. Solo nombres y metadatos."""
    sha = sha_git(raiz)
    plantilla = f"{raiz.name}/{app}@{sha}"
    registro, boveda = [], []
    mapa = defaultdict(lambda: defaultdict(set))  # (cat, base) -> ambiente -> llaves
    for ruta in buscar_archivos(app_dir):
        rel = ruta.relative_to(raiz).as_posix()
        amb, tok = ambiente_de(ruta)
        cat = "raíz" if ruta.parent.name in ("resources", app_dir.name) else ruta.parent.name
        base = base_de(ruta)
        for clave, valor in leer_llaves(ruta):
            sens = es_sensible(clave, valor)
            id_var = f"VAR-{id_inicio + len(registro):04d}"
            registro.append([id_var, app, amb, clave, "Sí" if sens else "No",
                             rel, plantilla, "v1", responsable, "Vigente"])
            if tok != "base":
                mapa[(cat, base)][tok].add(clave)
            if sens:
                ruta_boveda = f"kv/asp/{app}/{tok if tok != 'base' else 'todos'}/{clave}".lower()
                boveda.append([id_var, app, amb, clave, "[pendiente]", ruta_boveda, "v1", ""])
    diferencias = []
    for (cat, base), por_amb in sorted(mapa.items()):
        if len(por_amb) < 2:
            continue
        todas = set().union(*por_amb.values())
        for tok in sorted(por_amb):
            faltan = sorted(todas - por_amb[tok])
            if faltan:
                diferencias.append(f"{app}: {cat}/{base} [{tok}] le faltan: {', '.join(faltan)}")
    return registro, boveda, diferencias, plantilla


def escribir_csv(salida, registro, boveda):
    enc = "utf-8-sig"  # BOM para que Excel muestre bien los acentos
    with open(salida / "registro_variables.csv", "w", newline="", encoding=enc) as f:
        w = csv.writer(f)
        w.writerow(["ID", "Aplicación", "Ambiente", "Variable", "Sensible", "Ubicación del valor",
                    "Plantilla en Git (repo@versión)", "Versión del valor", "Responsable", "Estado"])
        w.writerows(registro)
    with open(salida / "boveda_secretos.csv", "w", newline="", encoding=enc) as f:
        w = csv.writer(f)
        w.writerow(["ID de variable", "Aplicación", "Ambiente", "Nombre del secreto",
                    "ID de credencial Jenkins", "Ruta futura en bóveda (Etapa 3)",
                    "Versión del secreto", "Fecha de última rotación"])
        w.writerows(boveda)


def extender_validaciones(ws, ultima):
    """Las listas desplegables llegan solo hasta la fila 1000 en la plantilla; se amplían."""
    for dv in ws.data_validations.dataValidation:
        nuevas = []
        for parte in str(dv.sqref).split():
            extremos = parte.split(":")
            extremos[-1] = re.sub(r"\d+$", str(ultima), extremos[-1])
            nuevas.append(":".join(extremos))
        dv.sqref = " ".join(nuevas)


def llenar_libro(xlsx_in, xlsx_out, registro, boveda):
    """Llena las pestañas del libro. Requiere openpyxl. No toca la Bitácora ni el Plan."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    font_captura = Font(name="Arial", size=10, color="0000FF")
    sin_relleno = PatternFill(fill_type=None)
    wb = load_workbook(xlsx_in)
    ultima = FILA_INICIAL + max(len(registro), len(boveda), 1) + 200

    def vaciar(ws, desde=5, columnas=10):
        """Borra datos previos (de corridas anteriores) para que no queden filas viejas mezcladas."""
        for fila in ws.iter_rows(min_row=desde, max_row=ws.max_row, min_col=1, max_col=columnas):
            for cell in fila:
                cell.value = None
                cell.fill = sin_relleno

    # Registro de variables: la fila 5 era la de ejemplo; se vacía todo y se empieza en la 6
    ws = wb["Registro de variables"]
    vaciar(ws)
    for i, fila in enumerate(registro):
        for c, val in enumerate(fila, start=1):
            cell = ws.cell(row=FILA_INICIAL + i, column=c, value=val if val != "" else None)
            cell.font = font_captura
            cell.fill = sin_relleno
    extender_validaciones(ws, ultima)

    # Bóveda: misma lógica; la fecha de rotación queda en blanco hasta que se capture
    ws = wb["Bóveda (Etapa 1)"]
    vaciar(ws)
    for i, fila in enumerate(boveda):
        r = FILA_INICIAL + i
        for c, val in enumerate(fila, start=1):
            cell = ws.cell(row=r, column=c, value=val if val != "" else None)
            cell.font = font_captura
            cell.fill = sin_relleno
        ws.cell(row=r, column=9, value=f'=IF(ISNUMBER(H{r}),TODAY()-H{r},"")')
        ws.cell(row=r, column=10, value=f'=IF(ISNUMBER(I{r}),IF(I{r}>90,"Rotar","Vigente"),"")')
    extender_validaciones(ws, ultima)

    # Resumen: rangos amplios para que no se corten al crecer el registro
    ws = wb["Resumen"]
    R = "'Registro de variables'!"
    B = "'Bóveda (Etapa 1)'!"
    formulas = [
        f"=COUNTA({R}A{FILA_INICIAL}:A{ULTIMA_FILA_FORMULAS})",
        f'=COUNTIF({R}E{FILA_INICIAL}:E{ULTIMA_FILA_FORMULAS},"Sí")',
        f'=COUNTIF({R}J{FILA_INICIAL}:J{ULTIMA_FILA_FORMULAS},"Vigente")',
        f'=COUNTIF({R}J{FILA_INICIAL}:J{ULTIMA_FILA_FORMULAS},"En cambio")',
        f"=COUNTA('Bitácora de versiones'!C{FILA_INICIAL}:C{ULTIMA_FILA_FORMULAS})",
        f'=COUNTIF({B}J{FILA_INICIAL}:J{ULTIMA_FILA_FORMULAS},"Rotar")',
    ]
    for i, f in enumerate(formulas, start=4):
        ws.cell(row=i, column=2, value=f)

    wb.save(xlsx_out)


def main():
    ap = argparse.ArgumentParser(description="Llena el control de variables desde un repo local (sin valores).")
    ap.add_argument("--repo", required=True, help="Carpeta del monorepo (con --todas) o de una aplicación")
    ap.add_argument("--todas", action="store_true",
                    help="Recorre cada subcarpeta del --repo como una aplicación distinta")
    ap.add_argument("--xlsx", default=None, help="Libro plantilla Control_Variables_Produccion.xlsx")
    ap.add_argument("--solo-csv", action="store_true", help="Generar solo CSV, sin tocar el libro")
    ap.add_argument("--responsable", default="[pendiente]", help="Responsable para cada fila")
    ap.add_argument("--salida", default="salida_variables", help="Carpeta de salida")
    args = ap.parse_args()

    raiz = Path(args.repo).resolve()
    if not raiz.is_dir():
        sys.exit(f"No existe la carpeta: {raiz}")

    if args.todas:
        apps = descubrir_apps(raiz)
        if not apps:
            sys.exit("No encontré subcarpetas con .properties o .env dentro del --repo.")
    else:
        apps = [raiz]

    registro, boveda, diferencias = [], [], []
    print(f"\nRepositorio: {raiz}  (commit {sha_git(raiz)})")
    print(f"Aplicaciones: {len(apps)}")
    for app_dir in apps:
        reg, bov, dif, _ = extraer_app(raiz, app_dir, app_dir.name, args.responsable, len(registro) + 1)
        registro += reg
        boveda += bov
        diferencias += dif
        print(f"  {app_dir.name}: {len(reg)} variables, {len(bov)} sensibles")

    if not registro:
        sys.exit("No se encontraron .properties ni .env en las aplicaciones.")

    salida = Path(args.salida)
    salida.mkdir(parents=True, exist_ok=True)
    escribir_csv(salida, registro, boveda)

    print(f"\nTotal: {len(registro)} variables  |  sensibles: {len(boveda)}")
    print("\nDiferencias entre ambientes:")
    print("\n".join(f"  {d}" for d in diferencias) if diferencias else "  Ninguna.")

    if args.solo_csv or not args.xlsx:
        print("\nCSV generados (pégalos en las pestañas):")
    else:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            print("\n[aviso] openpyxl no está instalado: se generaron solo los CSV.")
            print("        Para llenar el libro directo: py -m pip install openpyxl")
            args.xlsx = None
        if args.xlsx:
            xlsx_in = Path(args.xlsx)
            if not xlsx_in.is_file():
                sys.exit(f"No existe el libro: {xlsx_in}")
            xlsx_out = salida / "Control_Variables_Produccion_llenado.xlsx"
            llenar_libro(xlsx_in, xlsx_out, registro, boveda)
            print(f"\nLibro llenado (copia nueva, el original no se tocó): {xlsx_out.resolve()}")

    print(f"CSV en: {salida.resolve()}")
    print("Revisa la columna Sensible a mano: se marca por nombre de llave (p. ej. 'clave' puede ser de negocio).")


if __name__ == "__main__":
    main()
